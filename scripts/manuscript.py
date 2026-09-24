"""Private document worker. Requests and results are JSON; no shell interpolation."""
import copy
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path
from lxml import etree as ET

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
PARSER = ET.XMLParser(resolve_entities=False, no_network=True)
PLACEHOLDER = "Write this section here."


def xml(data):
    return ET.fromstring(data, parser=PARSER)


def read_package(path):
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > 256 * 1024 * 1024:
            raise ValueError("Document exceeds the expanded size limit.")
        parts = {name: archive.read(name) for name in archive.namelist()}
    if "word/document.xml" not in parts or "[Content_Types].xml" not in parts:
        raise ValueError("A valid DOCX document is required.")
    if any("vbaProject" in name for name in parts):
        raise ValueError("Macros are not allowed.")
    for name, data in parts.items():
        if name.endswith(".rels"):
            for rel in xml(data):
                if rel.get("TargetMode") == "External" and not rel.get("Type", "").endswith("/hyperlink"):
                    raise ValueError("Embed linked images and objects before submitting.")
    return parts


def write_package(parts, root, output):
    parts["word/document.xml"] = ET.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def element(name, **attrs):
    node = ET.Element("{%s}%s" % (W, name))
    for key, value in attrs.items():
        node.set("{%s}%s" % (W, key), str(value))
    return node


def paragraph(text, style=None):
    p = element("p")
    if style:
        props = element("pPr")
        props.append(element("pStyle", val=style))
        p.append(props)
    run = element("r")
    node = element("t")
    node.text = text
    run.append(node)
    p.append(run)
    return p


def text(node):
    return "".join(node.xpath(".//w:t/text()", namespaces=NS))


def controls(root):
    result = {}
    for sdt in root.xpath(".//w:sdt", namespaces=NS):
        tags = sdt.xpath("./w:sdtPr/w:tag/@w:val", namespaces=NS)
        if tags and tags[0].startswith("eprism:"):
            result.setdefault(tags[0][7:], []).append(sdt)
    return result


def section_control(section):
    sdt = element("sdt")
    props = element("sdtPr")
    props.append(element("tag", val="eprism:" + section["key"]))
    props.append(element("alias", val=section["label"]))
    props.append(element("lock", val="sdtLocked"))
    sdt.append(props)
    body = element("sdtContent")
    body.append(paragraph(section["label"], "Heading1"))
    body.append(paragraph(PLACEHOLDER))
    sdt.append(body)
    return sdt


def seed(args):
    parts = read_package(args["input"])
    root = xml(parts["word/document.xml"])
    body = root.find("w:body", NS)
    existing = controls(root)
    for section in args["sections"]:
        if section["key"] in existing:
            continue
        control = section_control(section)
        token = "${section:" + section["key"] + "}"
        anchor = next((p for p in body.findall("w:p", NS) if text(p).strip() == token), None)
        if anchor is not None:
            body.replace(anchor, control)
        else:
            sect = body.find("w:sectPr", NS)
            body.insert(list(body).index(sect) if sect is not None else len(body), control)
    # Metadata remains read-only; section controls alone are editable.
    for index, node in enumerate(list(body)):
        if node.tag in ("{%s}sectPr" % W, "{%s}sdt" % W):
            continue
        wrapper = element("sdt")
        props = element("sdtPr")
        props.append(element("tag", val="eprism-metadata:" + str(index)))
        props.append(element("lock", val="sdtContentLocked"))
        wrapper.append(props)
        content = element("sdtContent")
        position = list(body).index(node)
        body.remove(node)
        content.append(node)
        wrapper.append(content)
        body.insert(position, wrapper)
    write_package(parts, root, args["output"])
    return {"ok": True}


def migrate(args):
    # docxcompose remaps images, numbering and relationships before tagged sections move.
    from docx import Document
    from docxcompose.composer import Composer
    composer = Composer(Document(args["input"]))
    target_body = composer.doc.element.body
    original_count = len(target_body)
    composer.append(Document(args["parent"]))
    composer.save(args["output"])
    parts = read_package(args["output"])
    root = xml(parts["word/document.xml"])
    target_body = root.find("w:body", NS)
    original = list(target_body)[:original_count - 1]
    source = list(target_body)[original_count - 1:-1]
    old = {}
    for node in source:
        for key, values in controls(node).items():
            old.update({key: values[0]})
        if node.tag == "{%s}sdt" % W:
            tags = node.xpath("./w:sdtPr/w:tag/@w:val", namespaces=NS)
            if tags and tags[0].startswith("eprism:"):
                old[tags[0][7:]] = node
    # Maps a completed-template section key to the proposal-template key holding the same
    # content, wherever SubmissionTemplateRegistry renamed it across the two templates for the
    # same research_type (action_proposal -> action_completed, basic_proposal ->
    # basic_completed) — every other rich_text key is identical between the two, so only
    # genuine renames need an entry here.
    aliases = {"ethical_issues": "ethical_considerations", "introduction_and_rationale": "context_and_rationale"}
    for node in original:
        if node.tag != "{%s}sdt" % W:
            continue
        tags = node.xpath("./w:sdtPr/w:tag/@w:val", namespaces=NS)
        key = tags[0][7:] if tags else ""
        previous = old.get(key)
        if previous is None:
            previous = old.get(aliases.get(key, ""))
        if previous is not None:
            destination = node.find("w:sdtContent", NS)
            previous_content = previous.find("w:sdtContent", NS)
            for child in list(destination)[1:]:
                destination.remove(child)
            for child in list(previous_content)[1:]:
                destination.append(copy.deepcopy(child))
    for node in source:
        if node.getparent() is target_body:
            target_body.remove(node)
    write_package(parts, root, args["output"])
    return {"ok": True}


def defaults(parts, kind):
    if "word/styles.xml" not in parts:
        return {}
    styles = xml(parts["word/styles.xml"])
    node = styles.find("w:docDefaults/w:" + kind + "Default/w:" + kind, NS)
    return {ET.QName(child).localname: dict(child.attrib) for child in node} if node is not None else {}


def fonts(attributes, parts):
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    theme = xml(parts["word/theme/theme1.xml"]) if "word/theme/theme1.xml" in parts else None
    result = {}
    for face in ("ascii", "hAnsi"):
        literal = attributes.get("{%s}%s" % (W, face))
        themed = attributes.get("{%s}%sTheme" % (W, face))
        if literal:
            result[face] = literal.casefold()
        elif themed and theme is not None:
            branch = "majorFont" if themed.startswith("major") else "minorFont"
            latin = theme.find(".//{%s}%s/{%s}latin" % (a, branch, a))
            if latin is not None:
                result[face] = latin.get("typeface", "").casefold()
    return result


def properties(style, styles, kind, seen=None):
    if style is None:
        return {}
    seen = set() if seen is None else seen
    key = style.get("{%s}styleId" % W)
    if key in seen:
        return {}
    seen.add(key)
    parent = style.find("w:basedOn", NS)
    result = properties(styles.get(parent.get("{%s}val" % W)), styles, kind, seen) if parent is not None else {}
    props = style.find("w:" + kind, NS)
    if props is not None:
        for child in props:
            result.setdefault(ET.QName(child).localname, {}).update(child.attrib)
    return result


def validate(args):
    parts = read_package(args["input"])
    baseline = read_package(args["template"])
    root, template = xml(parts["word/document.xml"]), xml(baseline["word/document.xml"])
    found = controls(root)
    errors, warnings, sections = [], [], {}
    for definition in args["sections"]:
        key = definition["key"]
        matches = found.get(key, [])
        if len(matches) != 1:
            errors.append(definition["label"] + ": required section marker is missing or duplicated.")
            sections[key] = ""
            continue
        body = matches[0].find("w:sdtContent", NS)
        children = list(body) if body is not None else []
        if not children or text(children[0]).strip() != definition["label"]:
            errors.append(definition["label"] + ": restore the required heading.")
        value = " ".join(text(node) for node in children[1:]).replace(PLACEHOLDER, "").strip()
        has_image = any(node.xpath(".//w:drawing", namespaces=NS) for node in children[1:])
        sections[key] = value or ("[Image]" if has_image else "")
        if definition.get("required", True) and not sections[key]:
            errors.append(definition["label"] + ": enter section content.")
    for metadata in template.xpath("//w:sdt[w:sdtPr/w:tag[starts-with(@w:val, 'eprism-metadata:')]]", namespaces=NS):
        tag = metadata.xpath("./w:sdtPr/w:tag/@w:val", namespaces=NS)[0]
        matches = root.xpath("//w:sdt[w:sdtPr/w:tag/@w:val=$tag]", namespaces=NS, tag=tag)
        if len(matches) != 1 or text(matches[0]) != text(metadata):
            errors.append("Restore the research metadata from the selected template.")
    all_text = "".join(root.xpath("//w:t/text()", namespaces=NS))
    if re.search(r"\$\{[^}]+\}", all_text):
        errors.append("Replace all unresolved template placeholders.")
    if root.xpath("//w:ins | //w:del", namespaces=NS):
        errors.append("Accept or reject tracked changes before submission.")
    for name in ("pgSz", "pgMar"):
        expected = template.xpath("//w:sectPr/w:" + name, namespaces=NS)
        actual = root.xpath("//w:sectPr/w:" + name, namespaces=NS)
        if expected:
            keys = ("w", "h", "orient") if name == "pgSz" else ("top", "right", "bottom", "left", "header", "footer", "gutter")
            fallback = {"orient": "portrait", "gutter": "0"}
            expected_attrs = {key: expected[-1].get("{%s}%s" % (W, key), fallback.get(key)) for key in keys}
            if not actual or any(any(node.get("{%s}%s" % (W, key), fallback.get(key)) != value for key, value in expected_attrs.items()) for node in actual):
                errors.append("Page size or margins differ from the selected template (" + name + ").")
    if "word/styles.xml" in baseline and "word/styles.xml" in parts:
        original_styles = {s.get("{%s}styleId" % W): s for s in xml(baseline["word/styles.xml"]).findall("w:style", NS)}
        current_styles = {s.get("{%s}styleId" % W): s for s in xml(parts["word/styles.xml"]).findall("w:style", NS)}
        for key, matches in found.items():
            for control in matches:
                for p in control.xpath(".//w:p", namespaces=NS):
                    if not text(p).strip():
                        continue
                    ids = p.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS)
                    style_id = ids[0] if ids else "Normal"
                    if style_id not in original_styles:
                        warnings.append(key + ": custom style " + style_id + " requires visual review.")
                        continue
                    for kind, checks in (("pPr", ("spacing",)), ("rPr", ("rFonts", "sz"))):
                        expected = defaults(baseline, kind)
                        current = defaults(parts, kind)
                        for name, attrs in properties(original_styles.get(style_id), original_styles, kind).items():
                            expected.setdefault(name, {}).update(attrs)
                        for name, attrs in properties(current_styles.get(style_id), current_styles, kind).items():
                            current.setdefault(name, {}).update(attrs)
                        direct = p.find("w:pPr", NS) if kind == "pPr" else None
                        if direct is not None:
                            for child in direct:
                                current.setdefault(ET.QName(child).localname, {}).update(child.attrib)
                        compared = tuple(prop for prop in checks if prop != "rFonts")
                        font_mismatch = kind == "rPr" and fonts(expected.get("rFonts", {}), baseline) != fonts(current.get("rFonts", {}), parts)
                        if font_mismatch or any(expected.get(prop) and any(current.get(prop, {}).get(k) != v for k, v in expected[prop].items()) for prop in compared):
                            errors.append(key + ": restore template font, size or line spacing for " + style_id + ".")
                        if kind == "rPr":
                            for rpr in p.xpath("./w:r/w:rPr", namespaces=NS):
                                for prop in checks:
                                    override = rpr.find("w:" + prop, NS)
                                    if override is not None and expected.get(prop):
                                        changed = any(expected[prop].get(k) != v for k, v in override.attrib.items() if k in expected[prop])
                                        if prop == "rFonts":
                                            combined = dict(current.get(prop, {}))
                                            combined.update(override.attrib)
                                            changed = fonts(combined, parts) != fonts(expected[prop], baseline)
                                        if changed:
                                            errors.append(key + ": direct font or size overrides differ from the template.")
    return {"valid": not errors, "errors": list(dict.fromkeys(errors)), "warnings": list(dict.fromkeys(warnings)),
            "sections": sections, "text": "\n\n".join(sections.values()), "validator": "eprism-docx-1"}


CHAPTER_TOKEN_RE = re.compile(r"^\$\{([A-Za-z0-9_]+)\}$")


def chapter_token(key):
    return "${" + key + "}"


def chapter_marker(key, closing=False):
    # Tiny white text carrying the same [[section:<key>]] convention every other engine's
    # rendered PDF already uses (see SubmissionHtmlTemplateRenderer::buildScalars() and the
    # now-retired SubmissionDocxPdfMerger::stampChapterMarker()) — inserted as a genuine text
    # run here so pdf-review.js's wireChapterNav() keeps finding it via ordinary PDF text
    # extraction, unchanged, regardless of which pipeline produced the page. The closing
    # marker has no navigation purpose of its own — it exists purely so apply_formatting() can
    # later delimit exactly which paragraphs are this chapter's own imported content (to scope
    # body/heading formatting to it) without needing any state passed between the two
    # operations.
    p = element("p")
    run = element("r")
    rpr = element("rPr")
    rpr.append(element("sz", val=2))
    rpr.append(element("color", val="FFFFFF"))
    run.append(rpr)
    node = element("t")
    node.text = ("[[/section:" if closing else "[[section:") + key + "]]"
    run.append(node)
    p.append(run)
    return p


def strip_direct_typography(document):
    # Mirrors DocxStyleNormalizer::stripTypographyOverrides() exactly (same tags, same
    # pPr-direct-child guard for jc/spacing) — kept as the one place chapter typography is
    # normalized now that splicing happens here, so there is no separate PHP-side pass copying
    # the admin template's whole styles.xml over the chapter's own (see assemble_chapters()'s
    # own doc comment for why that wholesale copy was dropped).
    body = document.element.body
    for tag in ("rFonts", "sz", "szCs"):
        for node in list(body.iter("{%s}%s" % (W, tag))):
            node.getparent().remove(node)
    for tag in ("jc", "spacing"):
        for node in list(body.iter("{%s}%s" % (W, tag))):
            parent = node.getparent()
            if parent is not None and parent.tag == "{%s}pPr" % W:
                parent.remove(node)


def assemble_chapters(args):
    """
    Splices each rich_text chapter's own .docx body into the admin's already scalar/each-filled
    front-matter template, at that chapter's own `${<key>}` placeholder paragraph — replacing
    the old design where a chapter was never inserted into the template at all, only converted
    and appended as its own separate PDF page (see SubmissionDocxPdfMerger, now retired for the
    per-chapter engine). Uses docxcompose.Composer.insert() (not append()) at the placeholder's
    own body index, which is docxcompose's own supported arbitrary-position API — it already
    handles cross-package image/diagram/shape/footnote relationship remapping (de-duplicating
    identical images by content hash), style merging by *name* (so a chapter's own named style
    reuses the template's style with that name, and a genuinely different style is copied in
    rather than clobbered — see its class-level review below), and numbering remapping with a
    fresh restart per chapter for the first list of each style it encounters.

    Deliberately does NOT copy the template's whole styles.xml over the chapter first (the old
    DocxStyleNormalizer::normalize() behavior) — that would discard any chapter-specific named
    style (or a built-in list/table style) the template itself doesn't happen to define, which
    is exactly the "destroys chapter-specific structures" failure mode docxcompose's own
    per-name merge avoids. Only *direct* run/paragraph typography overrides (rFonts/sz/szCs/
    jc/spacing) are stripped from the chapter first, so its headings/body text fall through to
    whatever the template's own same-named style resolves to, without losing any style the
    template doesn't define at all.

    Validation is deliberately strict and named-by-key (matches DocxTemplateFiller's own
    "errors loud, name the key" convention): a chapter with content but no matching placeholder,
    more than one matching placeholder, a placeholder found inline or inside a table cell
    instead of its own standalone paragraph, or a placeholder-shaped paragraph left over that
    isn't one of the chapters this call was told about, all raise a specific, actionable error
    rather than silently appending elsewhere or leaving a literal "${...}" in the manuscript.

    @param args: {"template": path to the filled front-matter docx, "chapters": [{"key": str,
                  "path": path to that chapter's own docx, or None if never written}, ...],
                  "output": path to save the assembled docx to}
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml.section import CT_SectPr
    from docx.text.paragraph import Paragraph
    from docxcompose.composer import Composer

    document = Document(args["template"])
    body = document.element.body
    composer = Composer(document)
    provided_keys = {chapter["key"] for chapter in args["chapters"]}

    for chapter in args["chapters"]:
        key = chapter["key"]
        path = chapter.get("path")
        needle = chapter_token(key)

        matches = [p for p in document.paragraphs if p.text.strip() == needle]

        if not matches:
            if path is None:
                # Nothing written for this chapter yet, and no placeholder to fill either —
                # nothing to do (a required-but-empty chapter is a readiness-time concern, not
                # this operation's).
                continue

            inline = any(needle in p.text and p.text.strip() != needle for p in document.paragraphs)
            in_table = any(
                Paragraph(cell_p, None).text.strip() == needle
                for tbl in body.iter(qn("w:tbl"))
                for cell_p in tbl.iter(qn("w:p"))
            )
            if inline or in_table:
                where = "inside a table cell" if in_table else "inline with other text"
                raise ValueError(
                    "Chapter placeholder " + needle + " for \"" + key + "\" must occupy its own "
                    "paragraph — found " + where + " instead."
                )
            raise ValueError(
                "Document template is missing a placeholder for chapter \"" + key + "\" — add "
                + needle + " as its own paragraph."
            )

        if len(matches) > 1:
            raise ValueError(
                "Document template has " + str(len(matches)) + " placeholders for chapter \""
                + key + "\" (" + needle + ") — it must appear exactly once."
            )

        anchor = matches[0]
        index = list(body).index(anchor._p)

        if path is not None:
            chapter_doc = Document(path)
            strip_direct_typography(chapter_doc)
            # Composer.insert() silently skips a CT_SectPr element (see its own source) — the
            # count of what it actually inserts excludes that, so the closing marker below
            # lands exactly after the chapter's own content, not one short.
            inserted_count = sum(1 for el in chapter_doc.element.body if not isinstance(el, CT_SectPr))
            body.insert(index, chapter_marker(key))
            composer.insert(index + 1, chapter_doc)
            body.insert(index + 1 + inserted_count, chapter_marker(key, closing=True))

        anchor._p.getparent().remove(anchor._p)

    unknown = sorted(set(
        match.group(1)
        for p in document.paragraphs
        for match in [CHAPTER_TOKEN_RE.match(p.text.strip())]
        if match and match.group(1) not in provided_keys
    ))
    if unknown:
        raise ValueError(
            "Unrecognized chapter placeholder(s) in the document template: "
            + ", ".join(chapter_token(key) for key in unknown) + "."
        )

    document.save(args["output"])
    return {"ok": True}


PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
    "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
    "rPr", "sectPr", "pPrChange",
]
RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
    "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden",
    "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr",
    "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
]
CHAPTER_MARKER_RE = re.compile(r"^\[\[(/?)section:([A-Za-z0-9_]+)\]\]$")
JC_VALUES = {"left": "left", "right": "right", "center": "center", "justify": "both"}


def set_ordered(parent, tag, order, **attrs):
    # Get-or-create a child by tag, inserting it at the position the OOXML schema requires
    # relative to whichever of its schema-ordered siblings already exist — real consumers
    # (Word, ONLYOFFICE) are lenient about this in practice, but there is no reason to hand them
    # an out-of-schema-order document when producing one in order costs nothing extra. Reused for
    # both a paragraph's own pPr children and a run's rPr children by passing the matching order
    # list; setting the same attrs again on an existing node keeps every setter idempotent.
    node = parent.find("w:" + tag, NS)
    if node is None:
        node = element(tag)
        position = order.index(tag)
        target = len(parent)
        for index, existing in enumerate(parent):
            name = ET.QName(existing).localname
            if name in order and order.index(name) > position:
                target = index
                break
        parent.insert(target, node)
    for key, value in attrs.items():
        node.set("{%s}%s" % (W, key), str(value))
    return node


def ensure_pPr(p):
    pPr = p.find("w:pPr", NS)
    if pPr is None:
        pPr = element("pPr")
        p.insert(0, pPr)
    return pPr


def walk_style_chain(style, styles):
    seen = set()
    while style is not None:
        key = style.get("{%s}styleId" % W)
        if key in seen:
            return
        seen.add(key)
        yield style
        parent = style.find("w:basedOn", NS)
        style = styles.get(parent.get("{%s}val" % W)) if parent is not None else None


def resolve_outline_level(pPr, style, styles):
    # Structural only: a paragraph's own direct w:outlineLvl, or the first one found walking its
    # style's w:basedOn chain — never inferred from bold, caps, font size, or paragraph text, so
    # a chapter author's own "Heading 1"-styled paragraph is recognized the same way regardless
    # of what its style happens to be named or how it happens to look.
    if pPr is not None:
        node = pPr.find("w:outlineLvl", NS)
        if node is not None:
            return int(node.get("{%s}val" % W, "9"))
    for candidate in walk_style_chain(style, styles):
        style_pPr = candidate.find("w:pPr", NS)
        if style_pPr is not None:
            node = style_pPr.find("w:outlineLvl", NS)
            if node is not None:
                return int(node.get("{%s}val" % W, "9"))
    return None


def resolve_is_caption(style, styles):
    # A caption is only ever recognized via a resolved style literally named "Caption" (Word's
    # own built-in style used by Insert Caption, and what python-docx's own add_paragraph(style=
    # "Caption") produces) — never by a paragraph starting with the word "Figure" or similar.
    for candidate in walk_style_chain(style, styles):
        name = candidate.find("w:name", NS)
        if name is not None and name.get("{%s}val" % W, "").strip().casefold() == "caption":
            return True
    return False


def resolve_is_list(pPr, style, styles):
    if pPr is not None and pPr.find("w:numPr", NS) is not None:
        return True
    for candidate in walk_style_chain(style, styles):
        style_pPr = candidate.find("w:pPr", NS)
        if style_pPr is not None and style_pPr.find("w:numPr", NS) is not None:
            return True
    return False


def apply_run_props(p, font=None, size=None, bold=None, italic=None):
    if font is None and size is None and bold is None and italic is None:
        return
    pPr = ensure_pPr(p)
    targets = [set_ordered(pPr, "rPr", PPR_ORDER)]
    for r in p.xpath(".//w:r", namespaces=NS):
        rpr = r.find("w:rPr", NS)
        if rpr is None:
            rpr = element("rPr")
            r.insert(0, rpr)
        targets.append(rpr)
    for rpr in targets:
        if font:
            set_ordered(rpr, "rFonts", RPR_ORDER, ascii=font, hAnsi=font, cs=font)
        if size is not None:
            # HTML form numeric inputs may be persisted as JSON strings.
            half_points = int(round(float(size) * 2))
            set_ordered(rpr, "sz", RPR_ORDER, val=half_points)
            set_ordered(rpr, "szCs", RPR_ORDER, val=half_points)
        if bold is not None:
            set_ordered(rpr, "b", RPR_ORDER, val="1" if bold else "0")
            set_ordered(rpr, "bCs", RPR_ORDER, val="1" if bold else "0")
        if italic is not None:
            set_ordered(rpr, "i", RPR_ORDER, val="1" if italic else "0")
            set_ordered(rpr, "iCs", RPR_ORDER, val="1" if italic else "0")


def apply_paragraph_props(pPr, alignment=None, line_spacing=None, space_before=None,
                           space_after=None, first_line_indent=None, keep_with_next=None,
                           page_break_before=None):
    if alignment in JC_VALUES:
        set_ordered(pPr, "jc", PPR_ORDER, val=JC_VALUES[alignment])
    if line_spacing is not None or space_before is not None or space_after is not None:
        spacing_attrs = {}
        if line_spacing is not None:
            spacing_attrs["line"] = int(round(float(line_spacing) * 240))
            spacing_attrs["lineRule"] = "auto"
        if space_before is not None:
            spacing_attrs["before"] = int(round(float(space_before) * 20))
        if space_after is not None:
            spacing_attrs["after"] = int(round(float(space_after) * 20))
        set_ordered(pPr, "spacing", PPR_ORDER, **spacing_attrs)
    if first_line_indent is not None:
        set_ordered(pPr, "ind", PPR_ORDER, firstLine=int(round(float(first_line_indent) * 1440)))
    if keep_with_next is not None:
        set_ordered(pPr, "keepNext", PPR_ORDER, val="1" if keep_with_next else "0")
    if page_break_before is not None:
        set_ordered(pPr, "pageBreakBefore", PPR_ORDER, val="1" if page_break_before else "0")


def apply_formatting(args):
    """
    Applies the admin's `manuscript_format_options` policy (see the column's own migration
    comment) to an already-assembled manuscript docx — writing to a new output path, never
    touching the caller's input in place; the caller always passes a temporary assembled copy,
    never a researcher or admin source file. Runs after assemble_chapters() and before
    ONLYOFFICE's own docx->PDF conversion.

    Classification is entirely DOCX-structural, never heuristic (see resolve_outline_level(),
    resolve_is_caption(), resolve_is_list() above): never bold, caps, font size, or a paragraph's
    own text. Body/heading/table/caption rules are scoped to only the paragraphs and tables
    physically between assemble_chapters()'s own [[section:<key>]] / [[/section:<key>]] marker
    pair for each chapter — the admin's own front-matter/title-page content around the chapters
    is left completely untouched, other than the document's page margins (which are document-wide
    by nature, not chapter-scoped, and are applied to every section's own sectPr so a landscape
    section keeps its own orientation while still getting the configured margins).

    Every property is applied only when the policy sets it explicitly; anything absent is left
    exactly as assemble_chapters() produced it. Every setter overwrites its own element by value
    rather than accumulating a new one, so applying the same policy twice in a row (e.g.
    regenerating after only an unrelated chapter edit) produces byte-identical formatting rather
    than drifting.

    A body rule's first_line_indent is never applied to a resolved list paragraph (its indentation
    already comes from its numbering definition); every other body property still applies to list
    text. A table or caption category with "inherit": true (or simply absent) is skipped entirely,
    leaving that content exactly as the template/chapter itself formats it. Only paragraph/run
    text properties and page margins are ever touched — images, numbering definitions, table
    geometry (grid/spans/merges), headers/footers (a separate part, never opened here), fields,
    bookmarks and manual page breaks are never read or rewritten by this operation.

    @param args: {"input": path to the assembled docx, "options": the policy dict (falsy, or
                  without "enabled": true, means "leave the manuscript exactly as assembled"),
                  "output": path to save the formatted docx to}
    """
    policy = args.get("options") or {}
    if not policy.get("enabled"):
        shutil.copyfile(args["input"], args["output"])
        return {"ok": True}

    parts = read_package(args["input"])
    root = xml(parts["word/document.xml"])
    body = root.find("w:body", NS)
    styles_root = xml(parts["word/styles.xml"]) if "word/styles.xml" in parts else None
    styles = ({s.get("{%s}styleId" % W): s for s in styles_root.findall("w:style", NS)}
              if styles_root is not None else {})

    def resolve_style(p):
        pPr = p.find("w:pPr", NS)
        ids = pPr.xpath("./w:pStyle/@w:val", namespaces=NS) if pPr is not None else []
        return pPr, styles.get(ids[0] if ids else "Normal")

    body_cfg = policy.get("body")
    heading1_cfg = policy.get("heading1")
    heading23_cfg = policy.get("heading23")
    table_cfg = policy.get("table")
    caption_cfg = policy.get("caption")

    in_chapter = False
    for node in list(body):
        tag = ET.QName(node).localname

        if tag == "p":
            marker = CHAPTER_MARKER_RE.match(text(node).strip())
            if marker is not None:
                in_chapter = marker.group(1) == ""
                continue
            if not in_chapter:
                continue

            pPr, style = resolve_style(node)

            if resolve_is_caption(style, styles):
                if caption_cfg and not caption_cfg.get("inherit"):
                    apply_run_props(node, caption_cfg.get("font"), caption_cfg.get("size"),
                                     None, caption_cfg.get("italic"))
                    apply_paragraph_props(ensure_pPr(node), alignment=caption_cfg.get("alignment"))
                continue

            level = resolve_outline_level(pPr, style, styles)
            if level == 0:
                if heading1_cfg:
                    apply_run_props(node, heading1_cfg.get("font"), heading1_cfg.get("size"),
                                     heading1_cfg.get("bold"), None)
                    apply_paragraph_props(
                        ensure_pPr(node), alignment=heading1_cfg.get("alignment"),
                        space_before=heading1_cfg.get("space_before"),
                        space_after=heading1_cfg.get("space_after"),
                        keep_with_next=heading1_cfg.get("keep_with_next"),
                        page_break_before=heading1_cfg.get("page_break_before"),
                    )
                continue
            if level in (1, 2):
                if heading23_cfg:
                    apply_run_props(node, heading23_cfg.get("font"), heading23_cfg.get("size"),
                                     heading23_cfg.get("bold"), None)
                    apply_paragraph_props(
                        ensure_pPr(node), alignment=heading23_cfg.get("alignment"),
                        space_before=heading23_cfg.get("space_before"),
                        space_after=heading23_cfg.get("space_after"),
                        keep_with_next=heading23_cfg.get("keep_with_next"),
                        page_break_before=heading23_cfg.get("page_break_before"),
                    )
                continue
            if level is not None:
                # A deeper heading level (H4+) than this policy has a category for — left
                # untouched rather than guessed into either bucket.
                continue

            if body_cfg:
                is_list = resolve_is_list(pPr, style, styles)
                apply_run_props(node, body_cfg.get("font"), body_cfg.get("size"), None, None)
                apply_paragraph_props(
                    ensure_pPr(node), alignment=body_cfg.get("alignment"),
                    line_spacing=body_cfg.get("line_spacing"),
                    space_before=body_cfg.get("space_before"),
                    space_after=body_cfg.get("space_after"),
                    first_line_indent=None if is_list else body_cfg.get("first_line_indent"),
                )

        elif tag == "tbl":
            if in_chapter and table_cfg and not table_cfg.get("inherit"):
                for cell_p in node.iter("{%s}p" % W):
                    apply_run_props(cell_p, table_cfg.get("font"), table_cfg.get("size"), None, None)
                    apply_paragraph_props(ensure_pPr(cell_p), alignment=table_cfg.get("alignment"))

    page_cfg = policy.get("page") or {}
    margins = {
        "top": page_cfg.get("margin_top"), "right": page_cfg.get("margin_right"),
        "bottom": page_cfg.get("margin_bottom"), "left": page_cfg.get("margin_left"),
    }
    margin_twips = {key: int(round(float(value) * 1440)) for key, value in margins.items() if value is not None}
    if margin_twips:
        for sectPr in root.xpath("//w:sectPr", namespaces=NS):
            pgMar = sectPr.find("w:pgMar", NS)
            if pgMar is None:
                pgMar = element("pgMar")
                sectPr.append(pgMar)
            for key, value in margin_twips.items():
                pgMar.set("{%s}%s" % (W, key), str(value))

    write_package(parts, root, args["output"])
    return {"ok": True}


def pdf(args, encrypt=False):
    import pikepdf
    with pikepdf.open(args["input"]) as document:
        for attachment in args.get("attachments", []):
            with pikepdf.open(attachment) as other:
                document.pages.extend(other.pages)
        options = {}
        if encrypt:
            options["encryption"] = pikepdf.Encryption(
                owner=args["owner_password"], user=args.get("user_password", ""), R=6,
                allow=pikepdf.Permissions(extract=False, modify_annotation=False, modify_assembly=False,
                    modify_form=False, modify_other=False, print_lowres=True, print_highres=True))
        document.save(args["output"], **options)
    return {"ok": True}


if __name__ == "__main__":
    try:
        request = json.load(sys.stdin)
        operation = sys.argv[1]
        handlers = {"seed": seed, "validate": validate, "migrate": migrate,
                    "assemble_chapters": assemble_chapters, "apply_formatting": apply_formatting,
                    "merge_pdf": pdf, "encrypt_pdf": lambda args: pdf(args, True)}
        print(json.dumps(handlers[operation](request)))
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
