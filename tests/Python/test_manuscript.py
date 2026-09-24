import copy
import hashlib
import importlib.util
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image
import pikepdf

spec = importlib.util.spec_from_file_location("worker", Path(__file__).resolve().parents[2] / "scripts/manuscript.py")
worker = importlib.util.module_from_spec(spec)
spec.loader.exec_module(worker)


class ManuscriptWorkerTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.sections = [{"key": "context", "label": "Context", "required": True},
                         {"key": "references", "label": "References", "required": True}]
        template = Document()
        template.add_paragraph("Research title: Science")
        template.save(self.root / "raw.docx")
        worker.seed({"input": self.root / "raw.docx", "output": self.root / "template.docx", "sections": self.sections})

    def filled(self, output="filled.docx"):
        parts = worker.read_package(self.root / "template.docx")
        root = worker.xml(parts["word/document.xml"])
        for nodes in worker.controls(root).values():
            body = nodes[0].find("w:sdtContent", worker.NS)
            body.remove(body[-1])
            body.append(worker.paragraph("Actual research content and references."))
        worker.write_package(parts, root, self.root / output)
        return parts, root

    def validate(self, name="filled.docx"):
        return worker.validate({"input": self.root / name, "template": self.root / "template.docx", "sections": self.sections})

    def test_placeholders_do_not_count_as_research(self):
        result = self.validate("template.docx")
        self.assertFalse(result["valid"])
        self.assertEqual(result["sections"]["context"], "")

    def test_filled_sections_pass_and_metadata_is_protected(self):
        parts, root = self.filled()
        self.assertTrue(self.validate()["valid"], self.validate()["errors"])
        metadata = root.xpath("//w:sdt[w:sdtPr/w:tag[starts-with(@w:val, 'eprism-metadata:')]]", namespaces=worker.NS)[0]
        metadata.xpath(".//w:t", namespaces=worker.NS)[0].text = "Changed metadata"
        worker.write_package(parts, root, self.root / "filled.docx")
        self.assertFalse(self.validate()["valid"])

    def test_deleted_section_and_changed_margins_fail(self):
        parts, root = self.filled()
        control = worker.controls(root)["references"][0]
        control.getparent().remove(control)
        root.xpath("//w:pgMar", namespaces=worker.NS)[0].set("{%s}left" % worker.W, "42")
        worker.write_package(parts, root, self.root / "filled.docx")
        report = self.validate()
        self.assertTrue(any("missing" in error for error in report["errors"]))
        self.assertTrue(any("margins" in error for error in report["errors"]))

    def test_split_run_placeholder_is_detected(self):
        parts, root = self.filled()
        body = worker.controls(root)["context"][0].find("w:sdtContent", worker.NS)
        p = worker.paragraph("${unresolved")
        p.append(copy.deepcopy(worker.paragraph("}")[0]))
        body.append(p)
        worker.write_package(parts, root, self.root / "filled.docx")
        self.assertTrue(any("unresolved" in error for error in self.validate()["errors"]))

    def test_migration_preserves_sections_images_and_parent(self):
        self.filled()
        # Add an image and table inside the source section, with real package relationships.
        document = Document(self.root / "filled.docx")
        image = BytesIO()
        Image.new("RGB", (20, 20), "red").save(image, format="PNG")
        image.seek(0)
        document.add_picture(image)
        picture = document.element.body[-2]
        context = document.element.body.xpath(".//w:sdt[w:sdtPr/w:tag/@w:val='eprism:context']/w:sdtContent")[0]
        context.append(picture)
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Research table"
        context.append(table._tbl)
        document.save(self.root / "parent.docx")
        before = hashlib.sha256((self.root / "parent.docx").read_bytes()).hexdigest()
        worker.migrate({"input": self.root / "template.docx", "parent": self.root / "parent.docx", "output": self.root / "next.docx"})
        report = self.validate("next.docx")
        self.assertTrue(report["valid"], report["errors"])
        self.assertIn("Research table", report["sections"]["context"])
        package = worker.read_package(self.root / "next.docx")
        self.assertTrue(any(name.startswith("word/media/") for name in package))
        self.assertEqual(before, hashlib.sha256((self.root / "parent.docx").read_bytes()).hexdigest())

    def test_migration_carries_content_across_a_renamed_section_key(self):
        # basic_proposal's "context_and_rationale" becomes basic_completed's
        # "introduction_and_rationale" (and action_proposal's "ethical_considerations" becomes
        # action_completed's "ethical_issues") — SubmissionTemplateRegistry renames the key for
        # the same chapter across the two templates, so the migrate() alias map is the only thing
        # standing between "approved proposal" and "researcher's completed-research draft opens
        # with all their proposal content silently gone."
        proposal_sections = [{"key": "context_and_rationale", "label": "Context and Rationale", "required": True}]
        completed_sections = [{"key": "introduction_and_rationale", "label": "Introduction and Rationale", "required": True}]

        proposal_source = Document()
        proposal_source.add_paragraph("Proposal template")
        proposal_source.save(self.root / "proposal_raw.docx")
        worker.seed({"input": self.root / "proposal_raw.docx", "output": self.root / "proposal_template.docx", "sections": proposal_sections})

        parts = worker.read_package(self.root / "proposal_template.docx")
        root = worker.xml(parts["word/document.xml"])
        body = worker.controls(root)["context_and_rationale"][0].find("w:sdtContent", worker.NS)
        body.remove(body[-1])
        body.append(worker.paragraph("The researcher's actual rationale text."))
        worker.write_package(parts, root, self.root / "proposal_filled.docx")

        completed_source = Document()
        completed_source.add_paragraph("Completed template")
        completed_source.save(self.root / "completed_raw.docx")
        worker.seed({"input": self.root / "completed_raw.docx", "output": self.root / "completed_template.docx", "sections": completed_sections})

        worker.migrate({"input": self.root / "completed_template.docx",
                         "parent": self.root / "proposal_filled.docx", "output": self.root / "completed_next.docx"})

        report = worker.validate({"input": self.root / "completed_next.docx",
                                   "template": self.root / "completed_template.docx", "sections": completed_sections})
        self.assertIn("The researcher's actual rationale text.", report["sections"]["introduction_and_rationale"])

    def test_pdf_encryption_preserves_pages_and_requires_owner_for_full_permissions(self):
        with pikepdf.new() as pdf:
            pdf.add_blank_page(page_size=(300, 400))
            pdf.save(self.root / "review.pdf")
        worker.pdf({"input": self.root / "review.pdf", "output": self.root / "final.pdf", "owner_password": "a-long-secret"}, True)
        with pikepdf.open(self.root / "final.pdf") as pdf:
            self.assertTrue(pdf.is_encrypted)
            self.assertEqual(len(pdf.pages), 1)
            self.assertEqual(list(pdf.pages[0].MediaBox), [0, 0, 300, 400])
            self.assertFalse(pdf.allow.extract)
        with pikepdf.open(self.root / "final.pdf", password="a-long-secret") as pdf:
            self.assertTrue(pdf.owner_password_matched)


class AssembleChaptersTest(unittest.TestCase):
    """
    assemble_chapters() splices each rich_text chapter's own .docx into the admin's
    front-matter template at that chapter's own ${<key>} placeholder paragraph, via
    docxcompose.Composer.insert() at the placeholder's own body position (not append()) — the
    per-chapter engine's replacement for the old design that never inserted a chapter into the
    template at all, only converted and appended it as a separate PDF page.
    """

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def make_image(self, color):
        image = BytesIO()
        Image.new("RGB", (12, 12), color).save(image, format="PNG")
        image.seek(0)
        return image

    def make_template(self, body_paragraphs):
        """A template with a heading before/after each placeholder, matching the admin
        authoring convention this feature is for: the heading text lives in the template, the
        placeholder is the only thing on its own paragraph."""
        document = Document()
        for text in body_paragraphs:
            document.add_paragraph(text)
        path = self.root / "template.docx"
        document.save(path)
        return path

    def test_inserts_chapter_content_between_surrounding_template_paragraphs_in_order(self):
        template = self.make_template([
            "CHAPTER 1", "CONTEXT AND RATIONALE", "${context_and_rationale}",
            "CHAPTER 2", "RESEARCH QUESTIONS", "${research_questions}",
            "Generated footer line",
        ])

        chapter_one = Document()
        chapter_one.add_paragraph("First chapter body text.")
        chapter_one.save(self.root / "chapter_one.docx")

        chapter_two = Document()
        chapter_two.add_paragraph("Second chapter body text.")
        chapter_two.save(self.root / "chapter_two.docx")

        worker.assemble_chapters({
            "template": template,
            "chapters": [
                {"key": "context_and_rationale", "path": str(self.root / "chapter_one.docx")},
                {"key": "research_questions", "path": str(self.root / "chapter_two.docx")},
            ],
            "output": self.root / "assembled.docx",
        })

        result = Document(self.root / "assembled.docx")
        paragraphs = [p.text for p in result.paragraphs if p.text.strip() != ""]
        # Both placeholders are gone; both admin headings and both chapter bodies survive, in
        # template order, with each chapter's body landing strictly between its own heading and
        # the next chapter's heading.
        self.assertNotIn("${context_and_rationale}", paragraphs)
        self.assertNotIn("${research_questions}", paragraphs)
        context_idx = paragraphs.index("First chapter body text.")
        questions_idx = paragraphs.index("Second chapter body text.")
        self.assertLess(paragraphs.index("CONTEXT AND RATIONALE"), context_idx)
        self.assertLess(context_idx, paragraphs.index("CHAPTER 2"))
        self.assertLess(paragraphs.index("RESEARCH QUESTIONS"), questions_idx)
        self.assertLess(questions_idx, paragraphs.index("Generated footer line"))
        # The researcher never types the chapter title — it's already admin-authored content
        # that this operation must never touch, let alone delete.
        self.assertEqual(paragraphs.count("CONTEXT AND RATIONALE"), 1)

    def test_chapter_marker_is_embedded_as_tiny_white_text_immediately_before_the_chapter(self):
        # pdf-review.js's wireChapterNav() finds a chapter's own first page by searching every
        # PDF page's extracted text for this exact "[[section:<key>]]" marker (see
        # SubmissionHtmlTemplateRenderer::buildScalars() for the HTML-engine's identical
        # convention) — inserted here as a real run so it survives an ordinary docx->PDF
        # conversion as genuine, extractable page text, not something added at the PDF level.
        template = self.make_template(["CHAPTER 1", "${context_and_rationale}"])
        chapter = Document()
        chapter.add_paragraph("First chapter body text.")
        chapter.save(self.root / "chapter.docx")

        worker.assemble_chapters({
            "template": template,
            "chapters": [{"key": "context_and_rationale", "path": str(self.root / "chapter.docx")}],
            "output": self.root / "assembled.docx",
        })

        result = Document(self.root / "assembled.docx")
        paragraphs = [p.text for p in result.paragraphs if p.text.strip() != ""]
        marker_idx = paragraphs.index("[[section:context_and_rationale]]")
        self.assertEqual(paragraphs[marker_idx + 1], "First chapter body text.")

        marker_paragraph = next(p for p in result.paragraphs if p.text == "[[section:context_and_rationale]]")
        run = marker_paragraph.runs[0]
        self.assertEqual(run.font.size.pt, 1)
        self.assertEqual(str(run.font.color.rgb), "FFFFFF")

    def test_chapter_has_a_closing_marker_delimiting_its_own_content(self):
        # apply_formatting() (see scripts/manuscript.py) needs to know exactly which paragraphs
        # are this chapter's own imported content, to scope body/heading rules to it without
        # touching the admin's own front matter — the closing marker is what makes that
        # possible without passing any extra state between the two operations.
        template = self.make_template(["CHAPTER 1", "${context_and_rationale}", "CHAPTER 2"])
        chapter = Document()
        chapter.add_paragraph("First chapter body text.")
        chapter.add_paragraph("Second chapter paragraph.")
        chapter.save(self.root / "chapter.docx")

        worker.assemble_chapters({
            "template": template,
            "chapters": [{"key": "context_and_rationale", "path": str(self.root / "chapter.docx")}],
            "output": self.root / "assembled.docx",
        })

        result = Document(self.root / "assembled.docx")
        paragraphs = [p.text for p in result.paragraphs]
        start_idx = paragraphs.index("[[section:context_and_rationale]]")
        end_idx = paragraphs.index("[[/section:context_and_rationale]]")
        self.assertEqual(paragraphs[start_idx + 1:end_idx], ["First chapter body text.", "Second chapter paragraph."])
        self.assertEqual(paragraphs[end_idx + 1], "CHAPTER 2")

    def test_placeholder_split_across_multiple_runs_is_still_found(self):
        template = self.make_template(["CHAPTER 1"])
        document = Document(template)
        p = document.add_paragraph()
        p.add_run("${context_and")
        p.add_run("_rationale}")
        document.save(template)

        chapter = Document()
        chapter.add_paragraph("Body text.")
        chapter.save(self.root / "chapter.docx")

        worker.assemble_chapters({
            "template": template,
            "chapters": [{"key": "context_and_rationale", "path": str(self.root / "chapter.docx")}],
            "output": self.root / "assembled.docx",
        })

        result_text = [p.text for p in Document(self.root / "assembled.docx").paragraphs]
        self.assertIn("Body text.", result_text)
        self.assertNotIn("${context_and_rationale}", "".join(result_text))

    def test_a_chapter_with_no_content_yet_removes_the_placeholder_without_inserting_anything(self):
        template = self.make_template(["CHAPTER 1", "${context_and_rationale}", "CHAPTER 2"])

        worker.assemble_chapters({
            "template": template,
            "chapters": [{"key": "context_and_rationale", "path": None}],
            "output": self.root / "assembled.docx",
        })

        paragraphs = [p.text for p in Document(self.root / "assembled.docx").paragraphs]
        self.assertNotIn("${context_and_rationale}", paragraphs)
        self.assertEqual(paragraphs, ["CHAPTER 1", "CHAPTER 2"])

    def test_missing_placeholder_for_a_chapter_with_content_raises_a_named_error(self):
        template = self.make_template(["CHAPTER 1", "No placeholder here."])
        chapter = Document()
        chapter.add_paragraph("Body text.")
        chapter.save(self.root / "chapter.docx")

        with self.assertRaises(ValueError) as ctx:
            worker.assemble_chapters({
                "template": template,
                "chapters": [{"key": "context_and_rationale", "path": str(self.root / "chapter.docx")}],
                "output": self.root / "assembled.docx",
            })
        self.assertIn("context_and_rationale", str(ctx.exception))
        self.assertIn("missing", str(ctx.exception))

    def test_duplicate_placeholder_raises_a_named_error(self):
        template = self.make_template(["${context_and_rationale}", "CHAPTER 2", "${context_and_rationale}"])
        chapter = Document()
        chapter.add_paragraph("Body text.")
        chapter.save(self.root / "chapter.docx")

        with self.assertRaises(ValueError) as ctx:
            worker.assemble_chapters({
                "template": template,
                "chapters": [{"key": "context_and_rationale", "path": str(self.root / "chapter.docx")}],
                "output": self.root / "assembled.docx",
            })
        self.assertIn("context_and_rationale", str(ctx.exception))
        self.assertIn("exactly once", str(ctx.exception))

    def test_inline_placeholder_placement_raises_a_named_error(self):
        template = self.make_template(["See here: ${context_and_rationale} for details."])
        chapter = Document()
        chapter.add_paragraph("Body text.")
        chapter.save(self.root / "chapter.docx")

        with self.assertRaises(ValueError) as ctx:
            worker.assemble_chapters({
                "template": template,
                "chapters": [{"key": "context_and_rationale", "path": str(self.root / "chapter.docx")}],
                "output": self.root / "assembled.docx",
            })
        self.assertIn("own paragraph", str(ctx.exception))
        self.assertIn("inline", str(ctx.exception))

    def test_table_cell_placeholder_placement_raises_a_named_error(self):
        document = Document()
        document.add_paragraph("CHAPTER 1")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "${context_and_rationale}"
        template = self.root / "template.docx"
        document.save(template)

        chapter = Document()
        chapter.add_paragraph("Body text.")
        chapter.save(self.root / "chapter.docx")

        with self.assertRaises(ValueError) as ctx:
            worker.assemble_chapters({
                "template": template,
                "chapters": [{"key": "context_and_rationale", "path": str(self.root / "chapter.docx")}],
                "output": self.root / "assembled.docx",
            })
        self.assertIn("table cell", str(ctx.exception))

    def test_unrecognized_leftover_placeholder_raises_a_named_error(self):
        template = self.make_template(["${context_and_rationale}", "${a_typo_of_some_kind}"])
        chapter = Document()
        chapter.add_paragraph("Body text.")
        chapter.save(self.root / "chapter.docx")

        with self.assertRaises(ValueError) as ctx:
            worker.assemble_chapters({
                "template": template,
                "chapters": [{"key": "context_and_rationale", "path": str(self.root / "chapter.docx")}],
                "output": self.root / "assembled.docx",
            })
        self.assertIn("a_typo_of_some_kind", str(ctx.exception))

    def test_images_with_colliding_relationship_ids_across_chapters_both_survive(self):
        template = self.make_template(["${chapter_one}", "${chapter_two}"])

        chapter_one = Document()
        chapter_one.add_paragraph("Chapter one text.")
        chapter_one.add_picture(self.make_image("red"))
        chapter_one.save(self.root / "chapter_one.docx")

        chapter_two = Document()
        chapter_two.add_paragraph("Chapter two text.")
        chapter_two.add_picture(self.make_image("blue"))
        chapter_two.save(self.root / "chapter_two.docx")

        worker.assemble_chapters({
            "template": template,
            "chapters": [
                {"key": "chapter_one", "path": str(self.root / "chapter_one.docx")},
                {"key": "chapter_two", "path": str(self.root / "chapter_two.docx")},
            ],
            "output": self.root / "assembled.docx",
        })

        package = worker.read_package(self.root / "assembled.docx")
        media = [name for name in package if name.startswith("word/media/")]
        self.assertEqual(len(media), 2, "Expected both chapters' own distinct images to survive as separate media parts.")

        result = Document(self.root / "assembled.docx")
        # Every image reference must resolve to a real relationship on the assembled part —
        # a leftover/colliding rId here would raise a KeyError resolving the relationship.
        rids = result.element.body.xpath('.//a:blip/@r:embed')
        self.assertEqual(len(rids), 2)
        for rid in rids:
            self.assertIn(rid, result.part.rels)

    def test_numbered_and_nested_lists_restart_independently_per_chapter(self):
        template = self.make_template(["${chapter_one}", "${chapter_two}"])

        chapter_one = Document()
        chapter_one.add_paragraph("First.", style="List Number")
        chapter_one.add_paragraph("Second.", style="List Number")
        nested = chapter_one.add_paragraph("Nested.", style="List Number 2")
        chapter_one.save(self.root / "chapter_one.docx")

        chapter_two = Document()
        chapter_two.add_paragraph("Restarted first.", style="List Number")
        chapter_two.save(self.root / "chapter_two.docx")

        worker.assemble_chapters({
            "template": template,
            "chapters": [
                {"key": "chapter_one", "path": str(self.root / "chapter_one.docx")},
                {"key": "chapter_two", "path": str(self.root / "chapter_two.docx")},
            ],
            "output": self.root / "assembled.docx",
        })

        result = Document(self.root / "assembled.docx")
        texts = [p.text for p in result.paragraphs]
        self.assertIn("First.", texts)
        self.assertIn("Nested.", texts)
        self.assertIn("Restarted first.", texts)

        # docxcompose restarts numbering for the *first* paragraph of each style it encounters
        # per chapter (Composer.restart_first_numbering(), reset per insert() call) — so
        # "First." (chapter one) and "Restarted first." (chapter two) each get their own direct
        # <w:numPr> override with a distinct, really-defined numId, proving the two chapters'
        # lists were kept independent rather than sharing one running count.
        numbering = result.part.numbering_part.element
        num_ids_defined = set(numbering.xpath('.//w:num/@w:numId'))
        restarted = {p.text: p for p in result.paragraphs if p.text in ("First.", "Restarted first.")}
        self.assertEqual(set(restarted), {"First.", "Restarted first."})
        num_ids_used = []
        for text, p in restarted.items():
            num_id = p._p.xpath('.//w:numPr/w:numId/@w:val')
            self.assertTrue(num_id, "Expected a direct numId restart override on: %s" % text)
            self.assertIn(num_id[0], num_ids_defined)
            num_ids_used.append(num_id[0])
        self.assertEqual(len(set(num_ids_used)), 2, "Each chapter's restarted list must get its own distinct numId.")

    def test_tables_with_merged_cells_and_explicit_page_breaks_survive(self):
        from docx.enum.text import WD_BREAK

        template = self.make_template(["${chapter_one}"])

        chapter = Document()
        table = chapter.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "Merged header"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "B"
        chapter.add_paragraph("Before break.").add_run().add_break(WD_BREAK.PAGE)
        chapter.add_paragraph("After break.")
        chapter.save(self.root / "chapter.docx")

        worker.assemble_chapters({
            "template": template,
            "chapters": [{"key": "chapter_one", "path": str(self.root / "chapter.docx")}],
            "output": self.root / "assembled.docx",
        })

        result = Document(self.root / "assembled.docx")
        self.assertEqual(len(result.tables), 1)
        self.assertEqual(result.tables[0].cell(0, 0).text, "Merged header")
        self.assertTrue(result.tables[0].cell(0, 0)._tc.xpath('.//w:gridSpan'), "Expected the merged cell's gridSpan to survive.")
        self.assertTrue(result.element.body.xpath('.//w:br[@w:type="page"]'), "Expected the explicit page break to survive.")

    def test_conflicting_and_chapter_only_style_names_are_reconciled_by_name(self):
        from docx.enum.style import WD_STYLE_TYPE

        template = Document()
        template.styles.add_style("Chapter Highlight", WD_STYLE_TYPE.PARAGRAPH)
        template.add_paragraph("${chapter_one}")
        template_path = self.root / "template.docx"
        template.save(template_path)

        chapter = Document()
        # Same style *name* as the template already defines (but its own, independently-created
        # definition — a real cross-package "conflict" by id, reconciled by name) — must be
        # reused by name, not duplicated or left dangling.
        chapter.styles.add_style("Chapter Highlight", WD_STYLE_TYPE.PARAGRAPH)
        chapter.add_paragraph("Shared style text.", style="Chapter Highlight")
        # A style the template does NOT define at all — must be copied in, not dropped, so the
        # researcher's own formatting choice survives.
        chapter.styles.add_style("Researcher Only Style", WD_STYLE_TYPE.PARAGRAPH)
        chapter.add_paragraph("Chapter-only style text.", style="Researcher Only Style")
        chapter.add_paragraph("Ordinary subheading the researcher typed themselves.")
        chapter.save(self.root / "chapter.docx")

        worker.assemble_chapters({
            "template": template_path,
            "chapters": [{"key": "chapter_one", "path": str(self.root / "chapter.docx")}],
            "output": self.root / "assembled.docx",
        })

        result = Document(self.root / "assembled.docx")
        style_names = [s.name for s in result.styles]
        self.assertEqual(style_names.count("Chapter Highlight"), 1, "The shared-name style must not be duplicated.")
        self.assertIn("Researcher Only Style", style_names, "A chapter-only style must be copied in, not dropped.")
        texts = [p.text for p in result.paragraphs]
        self.assertIn("Shared style text.", texts)
        self.assertIn("Chapter-only style text.", texts)
        self.assertIn("Ordinary subheading the researcher typed themselves.", texts)

    def test_landscape_chapter_section_is_a_known_limitation_not_preserved(self):
        """
        Documents docxcompose's own acknowledged section-property limitation (see its README):
        Composer.insert()/append() drop an inserted document's own trailing (body-level)
        <w:sectPr> outright, so a chapter whose *only* section is landscape does not carry that
        orientation into the assembled manuscript — it inherits whatever orientation the
        template's own section already has. Not something this feature attempts to fix; the
        content itself (text/tables/images) still survives, only the page orientation doesn't.
        """
        from docx.enum.section import WD_ORIENT

        template = self.make_template(["${chapter_one}"])
        original_width, original_height = Document(template).sections[0].page_width, Document(template).sections[0].page_height

        chapter = Document()
        section = chapter.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        chapter.add_paragraph("Wide landscape content.")
        chapter.save(self.root / "chapter.docx")

        worker.assemble_chapters({
            "template": template,
            "chapters": [{"key": "chapter_one", "path": str(self.root / "chapter.docx")}],
            "output": self.root / "assembled.docx",
        })

        result = Document(self.root / "assembled.docx")
        texts = [p.text for p in result.paragraphs]
        self.assertIn("Wide landscape content.", texts, "The chapter's own content must still survive.")
        self.assertEqual(result.sections[0].page_width, original_width)
        self.assertEqual(result.sections[0].page_height, original_height)
        self.assertNotEqual(result.sections[0].orientation, WD_ORIENT.LANDSCAPE)

    def test_source_files_are_never_modified(self):
        template = self.make_template(["${chapter_one}"])
        chapter = Document()
        chapter.add_paragraph("Body text.")
        chapter_path = self.root / "chapter.docx"
        chapter.save(chapter_path)

        template_before = hashlib.sha256(Path(template).read_bytes()).hexdigest()
        chapter_before = hashlib.sha256(chapter_path.read_bytes()).hexdigest()

        worker.assemble_chapters({
            "template": template,
            "chapters": [{"key": "chapter_one", "path": str(chapter_path)}],
            "output": self.root / "assembled.docx",
        })

        self.assertEqual(template_before, hashlib.sha256(Path(template).read_bytes()).hexdigest())
        self.assertEqual(chapter_before, hashlib.sha256(chapter_path.read_bytes()).hexdigest())


class ApplyFormattingTest(unittest.TestCase):
    """
    apply_formatting() applies an admin's manuscript_format_options policy to an already-
    assembled manuscript docx (assemble_chapters()'s own output), between chapter assembly and
    ONLYOFFICE's own docx->PDF conversion. Classification is purely DOCX-structural (outline
    level / style name / numbering — never bold, caps, font size or paragraph text), and body/
    heading/table/caption rules apply only to the [[section:<key>]] / [[/section:<key>]] marker-
    delimited chapter content assemble_chapters() itself produces, never the admin's own
    surrounding front matter.
    """

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def make_image(self):
        image = BytesIO()
        Image.new("RGB", (12, 12), "red").save(image, format="PNG")
        image.seek(0)
        return image

    def assemble(self, template_paragraphs, build_chapter):
        template = Document()
        for para in template_paragraphs:
            if isinstance(para, tuple):
                template.add_paragraph(para[0], style=para[1])
            else:
                template.add_paragraph(para)
        template_path = self.root / "template.docx"
        template.save(template_path)

        chapter = Document()
        build_chapter(chapter)
        chapter_path = self.root / "chapter.docx"
        chapter.save(chapter_path)

        assembled_path = self.root / "assembled.docx"
        worker.assemble_chapters({
            "template": template_path,
            "chapters": [{"key": "chapter_one", "path": str(chapter_path)}],
            "output": assembled_path,
        })
        return assembled_path

    def format(self, input_path, options, output_name="formatted.docx"):
        output_path = self.root / output_name
        worker.apply_formatting({"input": input_path, "options": options, "output": output_path})
        return output_path

    def paragraph_by_text(self, doc, value):
        return next(p for p in doc.paragraphs if p.text == value)

    def test_disabled_or_absent_policy_copies_the_document_unchanged(self):
        assembled = self.assemble(["${chapter_one}"], lambda c: c.add_paragraph("Body text."))
        before = hashlib.sha256(Path(assembled).read_bytes()).hexdigest()

        for index, options in enumerate((None, {}, {"enabled": False})):
            output = self.format(assembled, options, "formatted_%d.docx" % index)
            self.assertEqual(hashlib.sha256(output.read_bytes()).hexdigest(), before)

    def test_numeric_strings_from_saved_form_match_numeric_formatting(self):
        def build(chapter):
            chapter.add_paragraph("Body text.")
            chapter.add_paragraph("Heading", style="Heading 1")
            chapter.add_paragraph("Subheading", style="Heading 2")
            chapter.add_paragraph("Caption text", style="Caption")
            chapter.add_table(rows=1, cols=1).cell(0, 0).text = "Table text"

        assembled = self.assemble(["Title page", "${chapter_one}"], build)
        numeric = {
            "enabled": True,
            "body": {"font": "Georgia", "alignment": "justify", "size": 12, "line_spacing": 1.5, "space_before": 0,
                     "space_after": 8, "first_line_indent": 0.5},
            "heading1": {"font": "Cambria", "alignment": "center", "bold": True,
                         "keep_with_next": True, "page_break_before": True,
                         "size": 16, "space_before": 12, "space_after": 6},
            "heading23": {"font": "Arial", "alignment": "left", "bold": False,
                          "keep_with_next": False, "page_break_before": False,
                          "size": 14, "space_before": 8, "space_after": 4},
            "table": {"inherit": False, "font": "Verdana", "alignment": "right", "size": 10},
            "caption": {"inherit": False, "font": "Georgia", "italic": True,
                        "alignment": "center", "size": 9},
            "page": {"margin_top": 1, "margin_bottom": 1,
                     "margin_left": 1.25, "margin_right": 1},
        }
        strings = copy.deepcopy(numeric)
        for group in strings.values():
            if isinstance(group, dict):
                for key, value in group.items():
                    if type(value) in (int, float):
                        group[key] = str(value)
        expected = self.format(assembled, numeric, "numeric.docx")
        actual = self.format(assembled, strings, "strings.docx")
        self.assertEqual(worker.read_package(expected)["word/document.xml"],
                         worker.read_package(actual)["word/document.xml"])

    def test_body_rules_apply_only_within_chapter_content(self):
        assembled = self.assemble(
            ["ADMIN HEADING", "${chapter_one}", "ADMIN FOOTER"],
            lambda c: c.add_paragraph("Researcher body text."),
        )
        output = self.format(assembled, {
            "enabled": True,
            "body": {"font": "Georgia", "size": 13, "alignment": "justify", "line_spacing": 1.5,
                     "space_after": 8, "first_line_indent": 0.5},
        })

        result = Document(output)
        chapter_p = self.paragraph_by_text(result, "Researcher body text.")
        run = chapter_p.runs[0]
        self.assertEqual(run.font.name, "Georgia")
        self.assertEqual(run.font.size.pt, 13)
        self.assertEqual(chapter_p.paragraph_format.alignment.name, "JUSTIFY")
        self.assertAlmostEqual(chapter_p.paragraph_format.first_line_indent.inches, 0.5)

        admin_heading = self.paragraph_by_text(result, "ADMIN HEADING")
        self.assertIsNone(admin_heading.runs[0].font.name)
        self.assertIsNone(admin_heading.paragraph_format.alignment)
        admin_footer = self.paragraph_by_text(result, "ADMIN FOOTER")
        self.assertIsNone(admin_footer.paragraph_format.alignment)

    def test_first_line_indent_is_never_applied_to_list_paragraphs(self):
        assembled = self.assemble(
            ["${chapter_one}"],
            lambda c: c.add_paragraph("List item.", style="List Number"),
        )
        output = self.format(assembled, {
            "enabled": True,
            "body": {"font": "Georgia", "first_line_indent": 0.5},
        })
        result = Document(output)
        item = self.paragraph_by_text(result, "List item.")
        self.assertEqual(item.runs[0].font.name, "Georgia")
        self.assertIsNone(item.paragraph_format.first_line_indent)

    def test_heading1_and_heading23_get_their_own_rules(self):
        def build(chapter):
            chapter.add_paragraph("Section Title", style="Heading 1")
            chapter.add_paragraph("Subsection Title", style="Heading 2")
            chapter.add_paragraph("Body under subsection.")

        assembled = self.assemble(["${chapter_one}"], build)
        output = self.format(assembled, {
            "enabled": True,
            "heading1": {"font": "Cambria", "size": 16, "bold": True, "alignment": "center"},
            "heading23": {"font": "Calibri", "size": 13, "bold": True, "alignment": "left"},
        })

        result = Document(output)
        h1 = self.paragraph_by_text(result, "Section Title")
        self.assertEqual(h1.runs[0].font.name, "Cambria")
        self.assertEqual(h1.runs[0].font.size.pt, 16)
        self.assertTrue(h1.runs[0].font.bold)
        self.assertEqual(h1.paragraph_format.alignment.name, "CENTER")

        h23 = self.paragraph_by_text(result, "Subsection Title")
        self.assertEqual(h23.runs[0].font.name, "Calibri")
        self.assertEqual(h23.runs[0].font.size.pt, 13)

    def test_heading_rules_do_not_touch_admin_front_matter_outside_chapter_markers(self):
        assembled = self.assemble(
            [("ADMIN TITLE", "Heading 1"), "${chapter_one}"],
            lambda c: c.add_paragraph("Chapter Heading", style="Heading 1"),
        )
        output = self.format(assembled, {
            "enabled": True,
            "heading1": {"font": "Cambria", "bold": True},
        })
        result = Document(output)
        admin_title = self.paragraph_by_text(result, "ADMIN TITLE")
        self.assertIsNone(admin_title.runs[0].font.name)
        chapter_heading = self.paragraph_by_text(result, "Chapter Heading")
        self.assertEqual(chapter_heading.runs[0].font.name, "Cambria")

    def test_heading_classification_uses_outline_level_not_style_name(self):
        # A custom style named nothing like "Heading" but structurally marked as outline level 0
        # (real Word behavior for a user-renamed heading style) must still be classified as
        # heading1 — proving classification is never a name/text heuristic.
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        def build(chapter):
            custom = chapter.styles.add_style("Custom Section Head", WD_STYLE_TYPE.PARAGRAPH)
            pPr = OxmlElement("w:pPr")
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), "0")
            pPr.append(outline)
            custom.element.append(pPr)
            chapter.add_paragraph("Non-Heading-Named Title", style="Custom Section Head")
            chapter.add_paragraph("Body text.")

        assembled = self.assemble(["${chapter_one}"], build)
        output = self.format(assembled, {
            "enabled": True,
            "heading1": {"font": "Cambria", "bold": True},
        })
        result = Document(output)
        custom_heading = self.paragraph_by_text(result, "Non-Heading-Named Title")
        self.assertEqual(custom_heading.runs[0].font.name, "Cambria")
        self.assertTrue(custom_heading.runs[0].font.bold)
        body_paragraph = self.paragraph_by_text(result, "Body text.")
        self.assertIsNone(body_paragraph.runs[0].font.name)

    def test_caption_style_is_recognized_by_style_only_not_by_text_heuristics(self):
        def build(chapter):
            chapter.add_paragraph("Figure 1: Not actually styled as a caption.")
            chapter.add_paragraph("Figure 2: A real caption.", style="Caption")

        assembled = self.assemble(["${chapter_one}"], build)
        output = self.format(assembled, {
            "enabled": True,
            "caption": {"font": "Consolas", "size": 9, "italic": True, "alignment": "center"},
        })
        result = Document(output)
        fake = self.paragraph_by_text(result, "Figure 1: Not actually styled as a caption.")
        self.assertIsNone(fake.runs[0].font.name)
        real = self.paragraph_by_text(result, "Figure 2: A real caption.")
        self.assertEqual(real.runs[0].font.name, "Consolas")
        self.assertTrue(real.runs[0].font.italic)
        self.assertEqual(real.paragraph_format.alignment.name, "CENTER")

    def test_table_cell_text_is_formatted_without_touching_geometry(self):
        def build(chapter):
            table = chapter.add_table(rows=2, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1))
            table.cell(0, 0).text = "Merged header"
            table.cell(1, 0).text = "A"
            table.cell(1, 1).text = "B"

        assembled = self.assemble(["${chapter_one}"], build)
        output = self.format(assembled, {
            "enabled": True,
            "table": {"font": "Verdana", "size": 10, "alignment": "center"},
        })
        result = Document(output)
        self.assertEqual(len(result.tables), 1)
        table = result.tables[0]
        self.assertTrue(table.cell(0, 0)._tc.xpath(".//w:gridSpan"), "Merge/geometry must survive.")
        run = table.cell(0, 0).paragraphs[0].runs[0]
        self.assertEqual(run.font.name, "Verdana")
        self.assertEqual(run.font.size.pt, 10)
        self.assertEqual(table.cell(0, 0).paragraphs[0].paragraph_format.alignment.name, "CENTER")

    def test_table_inherit_flag_skips_formatting_entirely(self):
        def build(chapter):
            table = chapter.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Untouched cell text."

        assembled = self.assemble(["${chapter_one}"], build)
        output = self.format(assembled, {
            "enabled": True,
            "table": {"inherit": True, "font": "Verdana"},
        })
        result = Document(output)
        run = result.tables[0].cell(0, 0).paragraphs[0].runs[0]
        self.assertIsNone(run.font.name)

    def test_page_margins_apply_to_every_section_preserving_orientation(self):
        from docx.enum.section import WD_ORIENT

        template = Document()
        template.add_paragraph("${chapter_one}")
        new_section = template.add_section()
        new_section.orientation = WD_ORIENT.LANDSCAPE
        new_section.page_width, new_section.page_height = new_section.page_height, new_section.page_width
        template_path = self.root / "template.docx"
        template.save(template_path)

        chapter = Document()
        chapter.add_paragraph("Body text.")
        chapter_path = self.root / "chapter.docx"
        chapter.save(chapter_path)

        assembled_path = self.root / "assembled.docx"
        worker.assemble_chapters({
            "template": template_path,
            "chapters": [{"key": "chapter_one", "path": str(chapter_path)}],
            "output": assembled_path,
        })

        output = self.format(assembled_path, {
            "enabled": True,
            "page": {"margin_top": 1.5, "margin_left": 1.25},
        })
        result = Document(output)
        self.assertEqual(len(result.sections), 2)
        self.assertEqual(result.sections[1].orientation, WD_ORIENT.LANDSCAPE)
        for section in result.sections:
            self.assertAlmostEqual(section.top_margin.inches, 1.5)
            self.assertAlmostEqual(section.left_margin.inches, 1.25)

    def test_applying_the_same_policy_twice_is_idempotent(self):
        def build(chapter):
            chapter.add_paragraph("Body text.")
            chapter.add_paragraph("A Title", style="Heading 1")

        assembled = self.assemble(["${chapter_one}"], build)
        options = {
            "enabled": True,
            "body": {"font": "Georgia", "size": 12, "alignment": "justify", "line_spacing": 1.5, "space_after": 8},
            "heading1": {"font": "Cambria", "size": 16, "bold": True, "alignment": "center", "keep_with_next": True},
            "page": {"margin_top": 1.25},
        }
        once = self.format(assembled, options, "once.docx")
        twice = self.format(once, options, "twice.docx")
        self.assertEqual(
            worker.read_package(once)["word/document.xml"],
            worker.read_package(twice)["word/document.xml"],
        )

    def test_images_lists_and_manual_page_breaks_survive_formatting(self):
        from docx.enum.text import WD_BREAK

        def build(chapter):
            chapter.add_paragraph("First.", style="List Number")
            chapter.add_picture(self.make_image())
            chapter.add_paragraph("Before break.").add_run().add_break(WD_BREAK.PAGE)
            chapter.add_paragraph("After break.")

        assembled = self.assemble(["${chapter_one}"], build)
        output = self.format(assembled, {
            "enabled": True,
            "body": {"font": "Georgia", "size": 12},
        })
        result = Document(output)
        self.assertTrue(result.element.body.xpath(".//a:blip"), "Image must survive.")
        list_item = self.paragraph_by_text(result, "First.")
        self.assertTrue(list_item._p.xpath(".//w:numPr"), "List numbering must survive.")
        self.assertTrue(result.element.body.xpath('.//w:br[@w:type="page"]'), "Manual page break must survive.")

    def test_input_file_is_never_modified(self):
        assembled = self.assemble(["${chapter_one}"], lambda c: c.add_paragraph("Body text."))
        before = hashlib.sha256(Path(assembled).read_bytes()).hexdigest()
        self.format(assembled, {"enabled": True, "body": {"font": "Georgia"}})
        self.assertEqual(hashlib.sha256(Path(assembled).read_bytes()).hexdigest(), before)


if __name__ == "__main__":
    unittest.main()
